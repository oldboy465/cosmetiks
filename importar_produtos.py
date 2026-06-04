import os
import sys
from datetime import datetime
from decimal import Decimal
import openpyxl

# Inserção do ecossistema do Flask para dentro do contexto de script standalone
from app import create_app, db
from app.models.produto import Produto
from app.models.categoria import Categoria
from app.models.marca import Marca
from app.models.fornecedor import Fornecedor

# CONFIGURAÇÃO DE ALVO: Defina o ID do revendedor dono desses produtos
REVENDEDOR_ID = 1 
CAMINHO_PLANILHA = r"C:\Users\noteboy\Desktop\add_product.xlsx"

def rodar_importacao():
    app = create_app()
    with app.app_context():
        print(f" Leitura iniciada da planilha atualizada: {CAMINHO_PLANILHA}")
        
        if not os.path.exists(CAMINHO_PLANILHA):
            print(f"❌ Erro crítico: O arquivo não foi localizado em {CAMINHO_PLANILHA}")
            return

        try:
            wb = openpyxl.load_workbook(CAMINHO_PLANILHA, data_only=True)
            ws = wb.active
        except Exception as e:
            print(f"❌ Falha ao abrir o arquivo Excel: {str(e)}")
            return

        # Garantir a existência de ao menos um fornecedor para o revendedor (Requisito de Chave Estrangeira)
        fornecedor_padrao = Fornecedor.query.filter_by(revendedor_id=REVENDEDOR_ID).first()
        if not fornecedor_padrao:
            fornecedor_padrao = Fornecedor(
                revendedor_id=REVENDEDOR_ID,
                nome="Fornecedor Principal",
                observacoes="Criado automaticamente na importacao de lote"
            )
            db.session.add(fornecedor_padrao)
            db.session.commit()
            print(" Fornecedor operacional padrão estabelecido com sucesso.")

        produtos_inseridos = 0
        produtos_atualizados = 0

        # Iteração a partir da linha 2 (pulando os cabeçalhos textuais)
        for row in range(2, ws.max_row + 1):
            codigo_barras_raw = ws.cell(row=row, column=1).value
            nome_produto = ws.cell(row=row, column=2).value
            preco_custo_raw = ws.cell(row=row, column=3).value
            categoria_raw = ws.cell(row=row, column=4).value
            preco_venda_raw = ws.cell(row=row, column=5).value
            data_compra_raw = ws.cell(row=row, column=6).value
            marca_raw = ws.cell(row=row, column=7).value

            # Ignora linhas complementares inteiramente vazias
            if not nome_produto:
                continue

            # Sanitização e normalização das strings coletadas
            nome_produto = str(nome_produto).strip()
            codigo_barras = str(codigo_barras_raw).strip() if codigo_barras_raw and str(codigo_barras_raw).strip() != "0" else None
            categoria_nome = str(categoria_raw).strip() if categoria_raw and str(categoria_raw).strip() != "A definir" else "Geral"
            marca_nome = str(marca_raw).strip() if marca_raw and str(marca_raw).strip() != "A definir" else "Geral"

            # 1. Resolução ou criação dinâmica de Categoria
            categoria = Categoria.query.filter_by(revendedor_id=REVENDEDOR_ID, nome=categoria_nome).first()
            if not categoria:
                categoria = Categoria(revendedor_id=REVENDEDOR_ID, nome=categoria_nome)
                db.session.add(categoria)
                db.session.commit()

            # 2. Resolução ou criação dinâmica de Marca
            marca = Marca.query.filter_by(revendedor_id=REVENDEDOR_ID, nome=marca_nome).first()
            if not marca:
                marca = Marca(revendedor_id=REVENDEDOR_ID, nome=marca_nome)
                db.session.add(marca)
                db.session.commit()

            # 3. Pesquisa de duplicidade para saber se o produto já existe na base
            produto_existente = None
            if codigo_barras:
                produto_existente = Produto.query.filter_by(revendedor_id=REVENDEDOR_ID, codigo_barras=codigo_barras).first()
            if not produto_existente:
                produto_existente = Produto.query.filter_by(revendedor_id=REVENDEDOR_ID, nome=nome_produto).first()

            if produto_existente:
                # LÓGICA ATUALIZADA: Se o produto já existe, conserta a Marca e a Categoria dele
                produto_existente.marca_id = marca.id
                produto_existente.categoria_id = categoria.id
                
                # Aproveita para certificar que o código de barras correto esteja salvo
                if codigo_barras and not produto_existente.codigo_barras:
                    produto_existente.codigo_barras = codigo_barras
                
                produtos_atualizados += 1
                continue

            # 4. Parsing seguro dos campos numéricos e monetários para novos produtos
            try:
                preco_custo = Decimal(str(preco_custo_raw or 0.00))
                preco_venda = Decimal(str(preco_venda_raw or 0.00))
            except Exception:
                preco_custo = Decimal('0.00')
                preco_venda = Decimal('0.00')

            # 5. Tratamento de data de compra
            data_compra = None
            if data_compra_raw:
                if isinstance(data_compra_raw, datetime):
                    data_compra = data_compra_raw.date()
                else:
                    try:
                        data_compra = datetime.strptime(str(data_compra_raw).strip(), '%Y-%m-%d').date()
                    except ValueError:
                        pass

            # 6. Construção e persistência do novo produto
            novo_produto = Produto(
                revendedor_id=REVENDEDOR_ID,
                categoria_id=categoria.id,
                marca_id=marca.id,
                fornecedor_id=fornecedor_padrao.id,
                nome=nome_produto,
                codigo_barras=codigo_barras,
                quantidade_estoque=0, 
                estoque_minimo=3,     
                preco_custo=preco_custo,
                preco_venda=preco_venda,
                data_compra=data_compra,
                situacao='Ativo'
            )
            db.session.add(novo_produto)
            produtos_inseridos += 1

        db.session.commit()
        print("\n=======================================================")
        print("  PROCESSAMENTO DE IMPORTAÇÃO E AJUSTE DE MARCAS CONCLUÍDO!")
        print(f"  Produtos novos inseridos com sucesso: {produtos_inseridos}")
        print(f"  Produtos antigos mapeados e corrigidos com a nova Marca: {produtos_atualizados}")
        print("=======================================================\n")

if __name__ == '__main__':
    rodar_importacao()